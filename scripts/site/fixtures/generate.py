"""Rebuild real sample files. See requirements.txt; then run render-assets.mjs."""
from pathlib import Path
import json
import os
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from docx import Document
from docx.shared import Pt as DocPt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
ROOT=Path(__file__).parent
scenes=json.loads((ROOT/'scenes.json').read_text())
pdfmetrics.registerFont(TTFont('ShowcaseCJK', os.environ.get('CJK_FONT', '/System/Library/Fonts/STHeiti Light.ttc'), subfontIndex=0))
for lang,c in scenes.items():
 out=ROOT/lang;out.mkdir(exist_ok=True);zh=lang=='zh'
 # PDF: an actual editorial creative brief, including selectable localized text.
 pdf=canvas.Canvas(str(out/'00-brief.pdf'),pagesize=(720,510));pdf.setFillColor(HexColor('#f1e5cf'));pdf.rect(0,0,720,510,fill=1,stroke=0)
 font='ShowcaseCJK' if zh else 'Helvetica'
 pdf.setFillColor(HexColor('#254a3b'));pdf.setFont('Helvetica-Bold',76);pdf.drawString(44,385,'EMBER')
 pdf.setFont(font,28);pdf.drawString(46,334,'慢慢醒来，好好喝一杯。' if zh else 'Slow mornings. Good coffee.')
 pdf.setFillColor(HexColor('#b84e2f'));pdf.circle(582,370,79,fill=1,stroke=0)
 pdf.setFillColor(HexColor('#f1e5cf'));pdf.setFont('Helvetica-Bold',19);pdf.drawCentredString(582,373,'FRESH');pdf.drawCentredString(582,350,'DAILY')
 pdf.setFillColor(HexColor('#254a3b'));pdf.setFont(font,15)
 lines=['品牌创意简报 / 2026 秋季','为街角的精品咖啡店，做一次温暖的开场。','交付：双语官网、三款包装、发布提案与预算。','纸张质感。深绿。陶土色。清晰的产品故事。'] if zh else ['CREATIVE BRIEF / AUTUMN 2026','A warm introduction to the neighborhood coffee roastery.','Deliver: bilingual site, packaging, launch deck and budget.','Paper. Forest green. Terracotta. A clear product story.']
 for i,line in enumerate(lines):pdf.drawString(46,222-i*35,line)
 pdf.setStrokeColor(HexColor('#254a3b'));pdf.line(46,260,674,260);pdf.save()
 # XLSX: usable budget data, formulas, and a real chart.
 wb=Workbook();ws=wb.active;ws.title='预算' if zh else 'Budget'
 headers=['项目','九月','十月','十一月','合计'] if zh else ['Launch budget','September','October','November','Total']
 ws.append(headers)
 labels=['包装设计','官网制作','店内物料','摄影','社交推广','预备金'] if zh else ['Packaging','Website','Print materials','Photography','Social campaign','Contingency']
 for i,label in enumerate(labels,2):ws.append([label,800+i*85,520+i*70,350+i*50,f'=SUM(B{i}:D{i})'])
 ws.append(['总计' if zh else 'TOTAL','=SUM(B2:B7)','=SUM(C2:C7)','=SUM(D2:D7)','=SUM(E2:E7)'])
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='254A3B');cell.font=Font(color='FFFFFF',bold=True,size=14)
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.font=Font(size=12);cell.alignment=Alignment(vertical='center')
 for col,w in [('A',25),('B',19),('C',19),('D',19),('E',19)]:ws.column_dimensions[col].width=w
 for i in range(1,9):ws.row_dimensions[i].height=30
 ws.freeze_panes='B2';chart=BarChart();chart.title='月度预算' if zh else 'Monthly budget';chart.add_data(Reference(ws,min_col=2,max_col=4,min_row=1,max_row=7),titles_from_data=True);chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=7));ws.add_chart(chart,'A11');wb.save(out/'01-budget.xlsx')
 # PPTX: actual editable slides rather than a screenshot renamed as a deck.
 deck=Presentation();deck.slide_width=Inches(12);deck.slide_height=Inches(8)
 for n in range(3):
  slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('B95032' if n==0 else 'F1E5CF')
  def text(x,y,w,h,value,size,color='F8EEDB'):
   box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));p=box.text_frame.paragraphs[0];run=p.add_run();run.text=value;run.font.size=Pt(size);run.font.name='Arial';run.font.bold=True;run.font.color.rgb=RGBColor.from_string(color)
  text(.65,.5,10,1,'EMBER / 2026',24)
  if n==0:
   text(.65,1.8,10,3,'为街角，\n烘焙好日子。' if zh else 'Roasted for\nthe everyday.',58)
   text(.7,6.4,10,1,'开店发布提案 · 品牌 / 包装 / 渠道' if zh else 'LAUNCH PRESENTATION · BRAND / PACKAGING / CHANNELS',19)
  else:text(.65,2,10,3,(['故事与用户','渠道与预算'][n-1] if zh else ['Story and audience','Channels and budget'][n-1]),44,'254A3B')
 deck.save(out/'02-launch.pptx')
 # DOCX: a compact, editable travel guide.
 doc=Document();doc.add_heading('京都，下过雨以后' if zh else 'Kyoto, after the rain',0)
 doc.add_paragraph('城市漫步 / 旅行手册' if zh else 'CITY WALKS / A FIELD GUIDE',style='Subtitle')
 for title,body in ([('从鸭川出发','沿着河边慢慢走，去一间独立书店避雨，再找一家只有六个座位的咖啡馆。'),('给下午留白','不赶景点。把一小时留给一条陌生小巷，把另一小时留给一本新书。'),('路线笔记','鸭川 → 出町柳 → 书店 → 咖啡馆。全程约三公里，留出三个小时。')] if zh else [('Start by the river','Follow the Kamo River, duck into an independent bookshop, then find a cafe with just six seats.'),('Leave the afternoon open','Skip the checklist. Leave one hour for an unfamiliar lane and another for a book you did not plan to buy.'),('Route notes','Kamo River → Demachiyanagi → bookshop → coffee. Around three kilometers; allow three hours.')]):
  doc.add_heading(title,1);doc.add_paragraph(body)
 doc.styles['Normal'].font.size=DocPt(12);doc.save(out/'03-guide.docx')
 # HTML: both a real interactive report and a visual design board.
 heading='阳台菜园 / 收成模拟' if zh else 'BALCONY GARDEN / HARVEST MODEL'
 report=f'''<!doctype html><html lang="{lang}"><meta charset="utf-8"><style>body{{margin:0;padding:38px;background:#e7ecda;color:#294532;font:18px Arial}}h1{{font-size:32px}}.metric{{font-size:70px;font-weight:bold}}svg{{width:100%}}label{{display:block;margin-top:18px}}</style><p>{heading}</p><h1>{'小小阳台，也能长出丰盛。' if zh else 'A little space. A generous harvest.'}</h1><span class="metric">8.4 kg</span><p>{'预计八周总收成 · 示例模型' if zh else 'Projected eight-week harvest · illustrative model'}</p><svg viewBox="0 0 640 220"><path d="M0 210L80 198L160 168L240 145L320 105L400 82L480 50L560 24L640 15V220H0Z" fill="#73956a"/><path d="M0 212L80 207L160 199L240 183L320 164L400 140L480 113L560 100L640 80" fill="none" stroke="#c17b47" stroke-width="6"/></svg><label>{'每天日照小时数' if zh else 'Hours of sunlight per day'} <input type="range" min="2" max="10" value="6" oninput="document.querySelector('.metric').textContent=(this.value*1.4).toFixed(1)+' kg'"></label></html>'''
 (out/'05-garden.html').write_text(report)
 name='慢慢醒来' if zh else 'Slow mornings'
 board=f'''<!doctype html><html lang="{lang}"><meta charset="utf-8"><style>*{{box-sizing:border-box}}body{{margin:0;background:#ede4d3;color:#244837;font-family:Arial}}main{{width:900px;height:620px;padding:40px;position:relative;overflow:hidden}}header{{letter-spacing:4px;font-size:14px}}h1{{font-family:Georgia;font-size:66px;margin:20px 0}}.bags{{display:flex;align-items:end;gap:25px;position:absolute;bottom:58px;left:110px;transform:rotate(-4deg)}}.bag{{width:185px;height:280px;background:#244837;box-shadow:13px 18px 0 #00000013;padding:35px 18px;color:#f4ead5;border-top:14px solid #19382a;position:relative}}.bag:nth-child(2){{background:#b84e31;border-color:#913b25;height:315px;transform:rotate(6deg)}}.bag:nth-child(3){{background:#d6ba83;border-color:#b19a69;color:#244837;transform:rotate(10deg)}}.brand{{font-weight:bold;font-size:34px;letter-spacing:-1px}}.line{{height:2px;background:currentColor;opacity:.6;margin:22px 0}}small{{position:absolute;bottom:20px;font-size:12px}}.circle{{position:absolute;right:55px;top:65px;width:115px;height:115px;border:2px solid #b84e31;border-radius:50%;display:grid;place-items:center;color:#b84e31;font-weight:bold;transform:rotate(14deg)}}</style><main><header>EMBER / {'咖啡包装设计' if zh else 'PACKAGING STUDY'}</header><h1>{name}.</h1><div class="circle">{'每日新鲜' if zh else 'FRESH DAILY'}</div><div class="bags">'''+''.join(f'<div class="bag"><div class="brand">EMBER</div><div class="line"></div><p>{label}</p><small>250 g / WHOLE BEAN</small></div>' for label in (['日常拼配','暖阳烘焙','山野单品'] if zh else ['EVERYDAY BLEND','SUNLIT ROAST','SINGLE ORIGIN']))+'</div></main></html>'
 (out/'packaging-source.html').write_text(board)
 # Side-by-side review board, deliberately sized for the real review-doc panel.
 review=f'''<!doctype html><html lang="{lang}"><meta charset="utf-8"><style>*{{box-sizing:border-box}}body{{margin:0;padding:22px;font:14px/1.5 Arial;color:#293d32;background:#f8f5ee}}h1{{font-size:22px;margin:0 0 18px}}.grid{{display:grid;grid-template-columns:1fr 1fr;gap:14px}}article{{border:1px solid #ddd4c3}}.cover{{height:225px;padding:24px 18px;background:#e7ddc8;position:relative}}.cover b{{font-size:38px;line-height:1.05;display:block;font-family:Georgia;margin-top:22px}}.cover i{{display:block;height:70px;width:42px;background:#244837;position:absolute;bottom:24px;right:25px;transform:rotate(-12deg)}}article:nth-child(2) .cover{{background:#b84e31;color:#fff0d6}}article:nth-child(2) .cover b{{font-family:Arial;font-size:38px}}article:nth-child(2) .cover i{{background:#e6c999;transform:rotate(12deg)}}.caption{{padding:15px}}.caption strong{{display:block}}footer{{margin-top:22px;padding-top:15px;border-top:1px solid #d8d3c7}}.palette{{display:flex;gap:7px;margin-top:14px}}.palette span{{width:34px;height:12px}}</style><h1>{c['reviewTitle']}</h1><div class="grid"><article><div class="cover">EMBER<b>{'慢慢<br>醒来。' if zh else 'Slow<br>mornings.'}</b><i></i></div><div class="caption"><strong>A / {'温暖的编辑感' if zh else 'Warm editorial'}</strong>{'纸张质感 · 日常故事' if zh else 'Paper textures · everyday stories'}<div class="palette"><span style="background:#244837"></span><span style="background:#d8c9aa"></span><span style="background:#b84e31"></span></div></div></article><article><div class="cover">EMBER<b>{'新鲜<br>出炉。' if zh else 'Freshly<br>roasted.'}</b><i></i></div><div class="caption"><strong>B / {'鲜明的图形感' if zh else 'Bold and graphic'}</strong>{'大字排版 · 货架辨识度' if zh else 'Bold type · shelf presence'}<div class="palette"><span style="background:#b84e31"></span><span style="background:#f3dfb4"></span><span style="background:#293c32"></span></div></div></article></div><footer>{'共同交付：三款包装、双语官网首屏、社交封面。<br>选定方向后，进入完整发布资料制作。' if zh else 'Both directions: three packages, a bilingual website hero and social covers.<br>Choose one direction for the full launch set.'}</footer></html>'''
 # Mirror the backend's iframe height handshake for this browser-only mock.
 review += '<script>new ResizeObserver(()=>parent.postMessage({__fleetAskHeight:document.body.scrollHeight},"*")).observe(document.body);</script>'
 (out/'review.html').write_text(review)
 (out/'07-checklist.md').write_text('# '+c['artifacts'][7]+'\n\n'+('\n'.join(['- [x] 核心购买路径验证','- [x] 中英文文案复核','- [x] 手机布局检查','- [x] 包装与官网素材导出','- [ ] 最终视觉方向确认','- [ ] 发布与回顾']) if zh else '\n'.join(['- [x] Core purchase paths checked','- [x] English and Chinese copy reviewed','- [x] Mobile layouts checked','- [x] Packaging and website assets exported','- [ ] Final visual direction approved','- [ ] Publish and review'])))
